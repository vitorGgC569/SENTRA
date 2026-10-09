"""Scoped file operations inspired by RPA Excel.Files, Tables and PDF keywords.

Real openpyxl/pypdf providers are optional. openpyxl never evaluates formulas;
only an explicitly configured Excel/LibreOffice worker may recalculate them.
CSV/TSV transformations use the standard library, not a Robot/Excel process.
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
import threading
import zipfile
import tempfile
import shutil
import sys
import uuid
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from .rpa import (AuthorizedPaths, DiagnosedExecutor, ExecutorFailure,
                  atomic_output, dependency, dependency_versions, digest,
                  effect_checkpoint, run_owned_worker, DOCUMENT_WORKER_SCRIPT)

DOCUMENT_ACTIONS = frozenset({"excel.inspect", "excel.transform", "table.inspect",
    "table.transform", "pdf.inspect", "pdf.text", "pdf.extract_pages",
    "artifact.verify", "acceptance.evaluate", "excel.recalculate", "pdf.ocr"})


@dataclass(frozen=True)
class FormulaRecalculationBinding:
    name: str
    owner_principal_id: str
    engine: str  # excel or libreoffice
    python_executable: str = sys.executable
    executable: str | None = None  # soffice, required for LibreOffice
    timeout_seconds: float = 60.0

    def __post_init__(self):
        if (not self.name or not self.owner_principal_id or self.engine not in {'excel','libreoffice'} or
                not Path(self.python_executable).is_absolute() or not 1 <= self.timeout_seconds <= 90 or
                self.engine == 'libreoffice' and (not self.executable or not Path(self.executable).is_absolute())):
            raise ValueError('invalid_recalculation_binding')


@dataclass(frozen=True)
class OCRBinding:
    name: str
    owner_principal_id: str
    tesseract_executable: str
    python_executable: str = sys.executable
    language: str = 'eng'
    dpi: int = 150
    max_pixels: int = 8_000_000
    timeout_seconds: float = 60.0

    def __post_init__(self):
        if (not self.name or not self.owner_principal_id or not Path(self.tesseract_executable).is_absolute() or
                not Path(self.python_executable).is_absolute() or not re.fullmatch(r'[A-Za-z0-9_+]+', self.language) or
                type(self.dpi) is not int or not 72 <= self.dpi <= 300 or
                type(self.max_pixels) is not int or not 1 <= self.max_pixels <= 32_000_000 or
                not 1 <= self.timeout_seconds <= 90):
            raise ValueError('invalid_ocr_binding')


@dataclass(frozen=True)
class DocumentBinding:
    capability_id: str
    paths: AuthorizedPaths
    actions: tuple[str, ...] = tuple(sorted(DOCUMENT_ACTIONS))
    timeout_seconds: float = 30.0
    max_input_bytes: int = 32 * 1024 * 1024
    max_expanded_bytes: int = 128 * 1024 * 1024
    max_rows: int = 10000
    max_columns: int = 128
    max_pages: int = 200
    max_text_chars: int = 65536
    recalculation_backends: tuple[FormulaRecalculationBinding, ...] = ()
    ocr_backends: tuple[OCRBinding, ...] = ()
    max_acceptance_criteria: int = 32

    def __post_init__(self):
        if (not self.capability_id or not self.actions or
                not set(self.actions) <= DOCUMENT_ACTIONS or
                not 1 <= self.timeout_seconds <= 120 or
                any(type(v) is not int or v <= 0 for v in (
                    self.max_input_bytes, self.max_expanded_bytes, self.max_rows,
                    self.max_columns, self.max_pages, self.max_text_chars))):
            raise ValueError("invalid_document_binding")
        for providers in (self.recalculation_backends, self.ocr_backends):
            if len({p.name for p in providers}) != len(providers) or any(
                    p.timeout_seconds + 5 >= self.timeout_seconds for p in providers):
                raise ValueError('provider_timeout_requires_larger_operation_budget')
        if not 1 <= self.max_acceptance_criteria <= 128:
            raise ValueError('invalid_acceptance_limit')


def compare_tables(actual_columns, actual_rows, expected_columns, expected_rows, *,
                   ignore_case=False, strip=False, ignore_column_order=False,
                   ignore_row_order=False, numeric_columns=(), key_columns=(),
                   absolute_tolerance=0.0, relative_tolerance=0.0, max_mismatches=20):
    """Deterministic artifact metric; duplicates count and tolerances are explicit.

    WindowsWorld compare_csv/table inspired comparison, with bounded mismatch
    evidence rather than a score silently hiding parser/provider failures.
    """
    if (type(absolute_tolerance) not in (int,float) or type(relative_tolerance) not in (int,float) or
            not math.isfinite(absolute_tolerance) or not math.isfinite(relative_tolerance) or
            absolute_tolerance < 0 or relative_tolerance < 0 or not 1 <= max_mismatches <= 100):
        raise ExecutorFailure('invalid_comparison_options')
    if (any(type(flag) is not bool for flag in (ignore_case,strip,ignore_column_order,ignore_row_order)) or
            not isinstance(numeric_columns,(list,tuple)) or not isinstance(key_columns,(list,tuple)) or
            any(type(c) is not str for c in tuple(numeric_columns)+tuple(key_columns))):
        raise ExecutorFailure('invalid_comparison_options')
    columns = list(expected_columns)
    if (set(actual_columns) != set(columns) or
            not ignore_column_order and list(actual_columns) != columns or
            not set(numeric_columns) <= set(columns) or not set(key_columns) <= set(columns)):
        return {'passed': False, 'score': 0.0, 'reason': 'columns_differ',
                'actual_columns': actual_columns, 'expected_columns': expected_columns}
    def normal(value, column):
        if column in numeric_columns and value is not None:
            try: value = float(value)
            except (TypeError,ValueError) as exc: raise ExecutorFailure('comparison_numeric_conversion_failed') from exc
            if not math.isfinite(value): raise ExecutorFailure('comparison_non_finite_value')
        if isinstance(value,str):
            if strip: value=value.strip()
            if ignore_case: value=value.casefold()
        return value
    actual = [[normal(row[c], c) for c in columns] for row in actual_rows]
    expected = [[normal(row[c], c) for c in columns] for row in expected_rows]
    if ignore_row_order:
        if numeric_columns and (absolute_tolerance or relative_tolerance) and not key_columns:
            raise ExecutorFailure('unordered_numeric_tolerance_requires_keys')
        def rowkey(row):
            selected = [row[columns.index(c)] for c in key_columns] if key_columns else row
            return json.dumps(selected, sort_keys=True, ensure_ascii=False, allow_nan=False)
        if key_columns and (len({rowkey(r) for r in actual}) != len(actual) or
                            len({rowkey(r) for r in expected}) != len(expected)):
            raise ExecutorFailure('comparison_duplicate_keys')
        actual.sort(key=rowkey); expected.sort(key=rowkey)
    mismatches=[]; total=abs(len(actual)-len(expected))
    for row_number, (left, right) in enumerate(zip(actual,expected),1):
        for column, a, e in zip(columns,left,right):
            match = (math.isclose(a,e,abs_tol=absolute_tolerance,rel_tol=relative_tolerance)
                if column in numeric_columns and a is not None and e is not None else
                type(a) is type(e) and a==e)
            if not match:
                total+=1
                if len(mismatches)<max_mismatches:
                    mismatches.append({'row':row_number,'column':column,
                        'actual':str(a)[:256],'expected':str(e)[:256]})
    passed=total==0
    return {'passed':passed,'score':1.0 if passed else 0.0,'mismatch_count':total,
            'mismatches':mismatches,'actual_rows':len(actual),'expected_rows':len(expected),
            'row_order_ignored':ignore_row_order}


def read_verifier_table(path: Path, binding: DocumentBinding, *, sheet=None, value_mode='formulas'):
    if path.stat().st_size > binding.max_input_bytes: raise ExecutorFailure('verifier_input_size_limit')
    if path.suffix.lower() == '.xlsx':
        provider=dependency('openpyxl',(3,1,5))
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist())>binding.max_expanded_bytes:
                raise ExecutorFailure('xlsx_expansion_limit')
        if value_mode not in {'formulas','cached_values'}: raise ExecutorFailure('invalid_verifier_value_mode')
        book=provider.load_workbook(path,read_only=True,data_only=value_mode=='cached_values',keep_links=False)
        try:
            ws=book[sheet] if sheet is not None else book.active
            if ws.max_row>binding.max_rows or ws.max_column>binding.max_columns:
                raise ExecutorFailure('verifier_dimension_limit')
            values=[[_json_cell(c.value) for c in row] for row in ws.iter_rows()]
            columns=[str(c) if c is not None else '' for c in values[0]] if values else []
            rows=[dict(zip(columns,row)) for row in values[1:]]
        finally: book.close()
    elif path.suffix.lower()=='.json':
        data=json.loads(path.read_text(encoding='utf-8'))
        if not isinstance(data,dict) or set(data)!={'columns','rows'}:
            raise ExecutorFailure('invalid_table_json')
        columns,rows=data['columns'],data['rows']
    elif path.suffix.lower() in {'.csv','.tsv'}:
        with path.open(encoding='utf-8-sig',newline='') as stream:
            reader=csv.DictReader(stream,delimiter='\t' if path.suffix.lower()=='.tsv' else ',')
            columns,rows=reader.fieldnames,[]
            for row in reader:
                rows.append(row)
                if len(rows)>binding.max_rows: raise ExecutorFailure('verifier_dimension_limit')
    else: raise ExecutorFailure('unsupported_verifier_format')
    if (not isinstance(columns,list) or not columns or any(type(c) is not str or not c for c in columns) or
            len(set(columns))!=len(columns) or len(columns)>binding.max_columns or not isinstance(rows,list) or
            len(rows)>binding.max_rows or any(not isinstance(r,dict) or set(r)!=set(columns) for r in rows)):
        raise ExecutorFailure('invalid_verifier_table')
    return columns,rows


def verify_artifact(binding: DocumentBinding, actual: str, expected: str, *, metric='table', options=None,
                    expected_sha256=None):
    """Read-only pure artifact comparison, no benchmark replay/postconfig."""
    effect_checkpoint()
    actual_path=binding.paths.resolve(actual); expected_path=binding.paths.resolve(expected)
    for path in (actual_path,expected_path):
        if path.stat().st_size>binding.max_input_bytes: raise ExecutorFailure('verifier_input_size_limit')
    expected_digest=digest(expected_path)
    if expected_sha256 is not None and expected_digest!=expected_sha256:
        raise ExecutorFailure('expected_artifact_digest_changed')
    options=options or {}
    if not isinstance(options,dict): raise ExecutorFailure('invalid_comparison_options')
    if metric=='bytes':
        if options: raise ExecutorFailure('invalid_comparison_options')
        passed=digest(actual_path)==expected_digest
        result={'passed':passed,'score':float(passed),'reason':'bytes_equal' if passed else 'bytes_differ'}
    elif metric=='table':
        allowed={'sheet','value_mode','ignore_case','strip','ignore_column_order','ignore_row_order',
                 'numeric_columns','key_columns','absolute_tolerance','relative_tolerance','max_mismatches'}
        if set(options)-allowed: raise ExecutorFailure('invalid_comparison_options')
        read_options={k:options[k] for k in ('sheet','value_mode') if k in options}
        left=read_verifier_table(actual_path,binding,**read_options)
        right=read_verifier_table(expected_path,binding,**read_options)
        result=compare_tables(*left,*right,**{k:v for k,v in options.items() if k not in read_options})
    elif metric=='pdf':
        if set(options)-{'compare_text','strip','ignore_case'}: raise ExecutorFailure('invalid_comparison_options')
        if any(type(v) is not bool for v in options.values()) or actual_path.suffix.lower()!='.pdf' or expected_path.suffix.lower()!='.pdf':
            raise ExecutorFailure('invalid_pdf_comparison_options')
        provider=dependency('pypdf',(6,10,0)); failures=[]
        with actual_path.open('rb') as af, expected_path.open('rb') as ef:
            a,e=provider.PdfReader(af,strict=True),provider.PdfReader(ef,strict=True)
            if a.is_encrypted or e.is_encrypted: raise ExecutorFailure('encrypted_pdf_unsupported')
            if max(len(a.pages),len(e.pages))>binding.max_pages: raise ExecutorFailure('pdf_page_limit')
            if len(a.pages)!=len(e.pages): failures.append('page_count')
            for i,(ap,ep) in enumerate(zip(a.pages,e.pages),1):
                if list(ap.mediabox)!=list(ep.mediabox): failures.append(f'page_{i}_dimensions')
                if options.get('compare_text',False):
                    at,et=ap.extract_text() or '',ep.extract_text() or ''
                    if max(len(at),len(et))>binding.max_text_chars: raise ExecutorFailure('pdf_text_verifier_limit')
                    if options.get('strip'): at,et=at.strip(),et.strip()
                    if options.get('ignore_case'): at,et=at.casefold(),et.casefold()
                    if at!=et: failures.append(f'page_{i}_text')
            result={'passed':not failures,'score':0.0 if failures else 1.0,'failures':failures[:20]}
    elif metric=='excel':
        result=compare_excel_artifacts(actual_path,expected_path,binding,**options)
    else: raise ExecutorFailure('unsupported_artifact_metric')
    if metric=='table' and '.xlsx' in {actual_path.suffix.lower(),expected_path.suffix.lower()}:
        result.update(formula_value_mode=options.get('value_mode','formulas'),
                      cached_values_may_be_missing_or_stale=options.get('value_mode')=='cached_values',recomputed=False)
    return dict(result,actual_sha256=digest(actual_path),expected_sha256=expected_digest,
                metric=metric,deterministic=True)


def compare_excel_artifacts(actual_path,expected_path,binding,*,sheets=None,value_mode='formulas',compare_styles=False):
    """Worksheet grids and optional public style properties, no header assumption."""
    if (Path(actual_path).suffix.lower()!='.xlsx' or Path(expected_path).suffix.lower()!='.xlsx' or
            value_mode not in {'formulas','cached_values'} or type(compare_styles) is not bool or
            sheets is not None and (not isinstance(sheets,list) or not sheets or
                any(type(s) is not str for s in sheets) or len(sheets)!=len(set(sheets)))):
        raise ExecutorFailure('invalid_excel_verifier_options')
    for path in (actual_path,expected_path):
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist())>binding.max_expanded_bytes:
                raise ExecutorFailure('xlsx_expansion_limit')
    lib=dependency('openpyxl',(3,1,5))
    actual=lib.load_workbook(actual_path,read_only=True,data_only=value_mode=='cached_values',keep_links=False)
    expected=None
    try:
        expected=lib.load_workbook(expected_path,read_only=True,data_only=value_mode=='cached_values',keep_links=False)
        names=sheets or expected.sheetnames
        if (len(names)>32 or not set(names)<=set(actual.sheetnames) or not set(names)<=set(expected.sheetnames) or
                sheets is None and actual.sheetnames!=expected.sheetnames):
            return {'passed':False,'score':0.0,'reason':'worksheet_names_differ'}
        differences=[]; count=0
        def color(c):
            return None if c is None else (c.type,str(c.rgb),str(c.indexed),str(c.theme),str(c.tint))
        def style(c):
            if not getattr(c,'has_style',False): return ('default',)
            return (c.number_format,c.font.name,c.font.sz,c.font.b,c.font.i,color(c.font.color),
                c.fill.patternType,color(c.fill.fgColor),c.alignment.horizontal,c.alignment.vertical,
                c.alignment.wrap_text,c.border.left.style,c.border.right.style,c.border.top.style,c.border.bottom.style)
        for name in names:
            effect_checkpoint()
            aw,ew=actual[name],expected[name]
            if max(aw.max_row,ew.max_row)>binding.max_rows or max(aw.max_column,ew.max_column)>binding.max_columns:
                raise ExecutorFailure('verifier_dimension_limit')
            if (aw.max_row,aw.max_column)!=(ew.max_row,ew.max_column):
                count+=1
                if len(differences)<20: differences.append({'sheet':name,'reason':'dimensions_differ'})
            for row_number,(ar,er) in enumerate(zip(aw.iter_rows(),ew.iter_rows()),1):
                for column_number,(a,e) in enumerate(zip(ar,er),1):
                    av,ev=_json_cell(a.value),_json_cell(e.value)
                    equal=(type(av) is type(ev) and av==ev) or (
                        type(av) in (int,float) and type(ev) in (int,float) and av==ev)
                    if not equal or compare_styles and style(a)!=style(e):
                        count+=1
                        if len(differences)<20: differences.append({'sheet':name,'row':row_number,
                            'column':column_number,'reason':'value_or_style_differ'})
        return {'passed':count==0,'score':float(count==0),'mismatch_count':count,'mismatches':differences,
                'value_mode':value_mode,'styles_compared':compare_styles,'recomputed':False,
                'cached_values_may_be_missing_or_stale':value_mode=='cached_values'}
    finally:
        actual.close()
        if expected is not None: expected.close()


def evaluate_acceptance(binding: DocumentBinding, criteria: list[dict], *, conjunction='and',
                        short_circuit=True, threshold=1.0):
    """OSWorld-inspired composites, with explicit skipped/error criteria.

    and/or use logical all/any. avg and sum expose scored aggregation. Unlike
    the upstream and average convention, 'and' here requires every criterion.
    """
    if (conjunction not in {'and','or','avg','sum'} or type(short_circuit) is not bool or
            type(threshold) not in (int,float) or not 0<threshold<=1 or
            not criteria or len(criteria)>binding.max_acceptance_criteria):
        raise ExecutorFailure('invalid_composite_acceptance')
    results=[]; stop=False
    if (any(not isinstance(c,dict) or type(c.get('id')) is not str or not c['id'] or
            set(c)-{'id','actual','expected','metric','options','expected_sha256'} for c in criteria) or
            len({c['id'] for c in criteria})!=len(criteria)):
        raise ExecutorFailure('invalid_composite_criterion')
    for c in criteria:
        # Even short-circuited criteria must name authorized real artifacts.
        binding.paths.resolve(c.get('actual')); binding.paths.resolve(c.get('expected'))
    for criterion in criteria:
        if stop:
            results.append({'id':criterion['id'],'state':'NOT_EVALUATED'}); continue
        result=verify_artifact(binding,criterion['actual'],criterion['expected'],
            metric=criterion.get('metric','table'),options=criterion.get('options'),
            expected_sha256=criterion.get('expected_sha256'))
        results.append(dict(result,id=criterion['id'],state='EVALUATED'))
        stop=short_circuit and ((conjunction=='and' and not result['passed']) or
                               (conjunction=='or' and result['passed']))
    evaluated=[r for r in results if r['state']=='EVALUATED']; scores=[r['score'] for r in evaluated]
    if conjunction=='and': score=float(all(r['passed'] for r in evaluated))
    elif conjunction=='or': score=float(any(r['passed'] for r in evaluated))
    elif conjunction=='sum': score=min(1.0,sum(scores))
    else: score=sum(scores)/len(scores)
    return {'passed':score>=threshold,'score':score,'conjunction':conjunction,'criteria':results,
            'deterministic':True,'short_circuited':stop and len(evaluated)<len(criteria)}


def _json_cell(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ExecutorFailure('non_finite_cell_value')
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    raise ExecutorFailure('unsupported_cell_type')


def transform_table(columns: list[str], rows: list[dict], *, conversions=None,
                    filters=None, select=None, sort_by=None, group_by=None, sums=None):
    """Pure ordered transforms: convert, filter, stable sort, group, project.

    Numeric conversion is explicit; unconverted CSV values compare as strings.
    Grouping retains first-seen order and uses deterministic count/sum columns.
    """
    result = [dict(row) for row in rows]
    def require(column):
        if column not in columns:
            raise ExecutorFailure('unknown_table_column')
    conversions = conversions or {}
    if not isinstance(conversions, dict):
        raise ExecutorFailure('invalid_conversions')
    for column, kind in conversions.items():
        require(column)
        if kind not in {'number', 'string'}:
            raise ExecutorFailure('unsupported_conversion')
        for row in result:
            try:
                value = float(row[column]) if kind == 'number' else str(row[column])
            except (ValueError, TypeError) as exc:
                raise ExecutorFailure('numeric_conversion_failed') from exc
            if isinstance(value, float) and not math.isfinite(value):
                raise ExecutorFailure('numeric_conversion_failed')
            row[column] = value
    for condition in filters or []:
        if not isinstance(condition, dict) or set(condition) != {'column', 'op', 'value'}:
            raise ExecutorFailure('invalid_filter')
        column, op, value = condition['column'], condition['op'], condition['value']
        require(column)
        if op not in {'eq', 'ne', 'gt', 'ge', 'lt', 'le', 'contains'}:
            raise ExecutorFailure('unsupported_filter')
        def matches(row):
            item = row[column]
            try:
                if op == 'eq': return item == value
                if op == 'ne': return item != value
                if op == 'gt': return item > value
                if op == 'ge': return item >= value
                if op == 'lt': return item < value
                if op == 'le': return item <= value
                return str(value) in str(item)
            except TypeError as exc:
                raise ExecutorFailure('incompatible_filter_types') from exc
        result = [row for row in result if matches(row)]
    for column in reversed(sort_by or []):
        require(column)
        try:
            result.sort(key=lambda row: (row[column] is None, row[column]))
        except TypeError as exc:
            raise ExecutorFailure('incompatible_sort_types') from exc
    output_columns = list(columns)
    if group_by:
        for column in group_by: require(column)
        for column in sums or []: require(column)
        output_columns = list(group_by) + ['count'] + [f'sum_{c}' for c in sums or []]
        if len(output_columns) != len(set(output_columns)):
            raise ExecutorFailure('group_column_collision')
        groups = {}
        for row in result:
            key = tuple(row[c] for c in group_by)
            try:
                group = groups.setdefault(key, dict(zip(group_by, key), count=0,
                    **{f'sum_{c}': 0.0 for c in sums or []}))
            except TypeError as exc:
                raise ExecutorFailure('invalid_group_key') from exc
            group['count'] += 1
            for column in sums or []:
                value = row[column]
                if type(value) not in (int, float) or not math.isfinite(value):
                    raise ExecutorFailure('sum_requires_numeric_conversion')
                group[f'sum_{column}'] += value
                if not math.isfinite(group[f'sum_{column}']):
                    raise ExecutorFailure('non_finite_aggregate')
        result = list(groups.values())
    elif sums:
        raise ExecutorFailure('sums_require_group_by')
    if select is not None:
        if not select or len(select) != len(set(select)) or not set(select) <= set(output_columns):
            raise ExecutorFailure('invalid_projection')
        output_columns = list(select)
        result = [{c: row[c] for c in output_columns} for row in result]
    return output_columns, result


class DocumentExecutor(DiagnosedExecutor):
    kind = 'documents'

    def __init__(self, *, machine_id, owner_principal_id, bindings, policy=None):
        if len({b.capability_id for b in bindings}) != len(bindings):
            raise ValueError('duplicate_document_capability')
        if any(p.owner_principal_id!=owner_principal_id for b in bindings
               for p in b.recalculation_backends+b.ocr_backends):
            raise ValueError('document_provider_owner_mismatch')
        super().__init__(machine_id=machine_id, owner_principal_id=owner_principal_id,
            bindings={b.capability_id: b for b in bindings}, policy=policy)
        self._file_lock = threading.Lock()

    def _validate(self, request, binding):
        a = dict(request.arguments)
        action = a.get('action')
        if not isinstance(action, str) or action not in binding.actions:
            raise ValueError('document_action_denied')
        if action in {'artifact.verify','acceptance.evaluate'}:
            return self._validate_verification(request,binding,a)
        fields = {
            'excel.inspect': {'sheet', 'value_mode', 'start_row', 'end_row'},
            'excel.transform': {'sheet', 'cells', 'output', 'overwrite'},
            'table.inspect': {'limit'},
            'table.transform': {'output', 'overwrite', 'conversions', 'filters',
                                'select', 'sort_by', 'group_by', 'sums'},
            'pdf.inspect': set(), 'pdf.text': {'pages'},
            'pdf.extract_pages': {'pages', 'output', 'overwrite'},
            'excel.recalculate': {'backend','output','overwrite','expected_cache'},
            'pdf.ocr': {'backend','output','overwrite','pages'},
        }[action]
        if set(a) - ({'action', 'input'} | fields) or 'input' not in a:
            raise ValueError('invalid_document_arguments')
        source = binding.paths.resolve(a['input'])
        if source.stat().st_size > binding.max_input_bytes:
            raise ValueError('input_size_limit')
        suffixes = { 'excel': {'.xlsx'}, 'table': {'.csv', '.tsv'}, 'pdf': {'.pdf'} }
        family = action.split('.')[0]
        if source.suffix.lower() not in suffixes[family]:
            raise ValueError('unsupported_input_format')
        a['input'] = str(source)
        if action.endswith(('transform', 'extract_pages','recalculate','ocr')):
            target = binding.paths.resolve(a.get('output'), write=True)
            if target == source or (target.exists() and target.samefile(source)):
                raise ValueError('source_overwrite_denied')
            if target.suffix.lower() not in ({'.json'} if action=='pdf.ocr' else
                    suffixes[family] | ({'.json'} if family == 'table' else set())):
                raise ValueError('unsupported_output_format')
            a['output'] = str(target)
            if type(a.get('overwrite', False)) is not bool:
                raise ValueError('invalid_overwrite')
        if 'sheet' in a and (type(a['sheet']) is not str or not a['sheet']):
            raise ValueError('invalid_sheet')
        if action == 'excel.inspect':
            if a.get('value_mode', 'formulas') not in {'formulas', 'cached_values'}:
                raise ValueError('explicit_formula_mode_required')
            for key in ('start_row', 'end_row'):
                if key in a and (type(a[key]) is not int or not 1 <= a[key] <= binding.max_rows):
                    raise ValueError('row_limit')
        if action == 'excel.transform':
            cells = a.get('cells')
            if not isinstance(cells, dict) or not cells or len(cells) > 1000:
                raise ValueError('invalid_cell_patches')
            for cell, value in cells.items():
                if not isinstance(cell, str) or not re.fullmatch('[A-Z]{1,3}[1-9][0-9]{0,6}', cell):
                    raise ValueError('invalid_cell_reference')
                if type(value) not in (str, int, float, bool, type(None)):
                    raise ValueError('invalid_cell_value')
                if isinstance(value, str) and (len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value)):
                    raise ValueError('invalid_cell_text')
                if type(value) is float and not math.isfinite(value):
                    raise ValueError('non_finite_value')
        for key in ('pages', 'select', 'sort_by', 'group_by', 'sums', 'filters'):
            if key in a and (not isinstance(a[key], list) or len(a[key]) > binding.max_rows):
                raise ValueError('invalid_list_argument')
        for key in ('select', 'sort_by', 'group_by', 'sums'):
            if key in a and (any(type(c) is not str for c in a[key]) or len(a[key]) != len(set(a[key]))):
                raise ValueError('invalid_column_list')
        if 'pages' in a and (not a['pages'] or any(type(p) is not int or p < 1 for p in a['pages'])):
            raise ValueError('pages_are_one_based')
        if 'limit' in a and (type(a['limit']) is not int or not 1 <= a['limit'] <= binding.max_rows):
            raise ValueError('invalid_table_limit')
        if action=='excel.recalculate':
            if a.get('backend') not in {p.name for p in binding.recalculation_backends}:
                raise ValueError('configured_recalculation_backend_required')
            expected=a.get('expected_cache')
            if not isinstance(expected,dict) or not expected or len(expected)>1000:
                raise ValueError('explicit_expected_formula_cache_required')
            for coordinate,value in expected.items():
                if (not isinstance(coordinate,str) or not re.fullmatch(r'[^!]+![A-Z]{1,3}[1-9][0-9]{0,6}',coordinate) or
                        type(value) not in (str,int,float,bool,type(None))):
                    raise ValueError('invalid_expected_formula_cache')
                if type(value) is float and not math.isfinite(value):
                    raise ValueError('invalid_expected_formula_cache')
        if action=='pdf.ocr' and a.get('backend') not in {p.name for p in binding.ocr_backends}:
            raise ValueError('configured_ocr_backend_required')
        a['_operation_id'] = request.operation_id
        return a

    def _validate_verification(self,request,b,a):
        fields=({'action','input','expected','metric','options','expected_sha256','output','overwrite'}
            if a['action']=='artifact.verify' else
            {'action','criteria','conjunction','short_circuit','threshold','output','overwrite'})
        if set(a)-fields: raise ValueError('invalid_verifier_arguments')
        criteria=([dict(actual=a.get('input'),expected=a.get('expected'),id='artifact',
            **{k:a[k] for k in ('metric','options','expected_sha256') if k in a})]
            if a['action']=='artifact.verify' else a.get('criteria'))
        if not isinstance(criteria,list) or not criteria or len(criteria)>b.max_acceptance_criteria:
            raise ValueError('invalid_acceptance_criteria')
        ids=[]
        for c in criteria:
            if (not isinstance(c,dict) or set(c)-{'id','actual','expected','metric','options','expected_sha256'} or
                    type(c.get('id')) is not str or not c['id'] or len(c['id'])>128):
                raise ValueError('invalid_acceptance_criterion')
            ids.append(c['id'])
            if c.get('metric','table') not in {'table','bytes','pdf','excel'}: raise ValueError('unsupported_verifier_metric')
            if 'options' in c and not isinstance(c['options'],dict): raise ValueError('invalid_verifier_options')
            if 'expected_sha256' in c and (type(c['expected_sha256']) is not str or
                    not re.fullmatch('[a-f0-9]{64}',c['expected_sha256'])):
                raise ValueError('invalid_expected_digest')
            for k in ('actual','expected'):
                path=b.paths.resolve(c.get(k))
                if path.stat().st_size>b.max_input_bytes: raise ValueError('verifier_input_size_limit')
        if len(set(ids))!=len(ids): raise ValueError('duplicate_acceptance_ids')
        if a['action']=='acceptance.evaluate':
            if (a.get('conjunction','and') not in {'and','or','avg','sum'} or
                type(a.get('short_circuit',True)) is not bool or
                type(a.get('threshold',1.0)) not in (int,float) or not 0<a.get('threshold',1.0)<=1):
                raise ValueError('invalid_composite_acceptance')
        if 'output' in a:
            target=b.paths.resolve(a['output'],write=True)
            if target.suffix.lower()!='.json' or any(target==b.paths.resolve(c[k]) or
                    target.exists() and target.samefile(b.paths.resolve(c[k]))
                    for c in criteria for k in ('actual','expected')):
                raise ValueError('acceptance_report_overwrites_input')
            if type(a.get('overwrite',False)) is not bool: raise ValueError('invalid_overwrite')
            a['output']=str(target)
        a['_operation_id']=request.operation_id
        return a

    def _run(self, b, a):
        # Serialize document publication inside this adapter; other processes
        # still require host fencing / immutable input ownership.
        with self._file_lock:
            if a['action'] in {'artifact.verify','acceptance.evaluate'}:
                if a['action']=='artifact.verify':
                    result=verify_artifact(b,a['input'],a['expected'],**{k:a[k] for k in
                        ('metric','options','expected_sha256') if k in a})
                else:
                    result=evaluate_acceptance(b,a['criteria'],**{k:a[k] for k in
                        ('conjunction','short_circuit','threshold') if k in a})
                payload=json.dumps(result,allow_nan=False,ensure_ascii=False).encode('utf-8')
                if len(payload)>b.max_text_chars: raise ExecutorFailure('acceptance_evidence_size_limit')
                if 'output' in a:
                    result['artifact']=atomic_output(b.paths,a['output'],lambda p:p.write_bytes(payload),
                        lambda p:json.loads(p.read_text(encoding='utf-8')),overwrite=a.get('overwrite',False))
                return result
            source = b.paths.resolve(a['input'])
            if source.stat().st_size > b.max_input_bytes:
                raise ExecutorFailure('input_size_limit')
            original = digest(source)
            if a['action'] in {'excel.recalculate','pdf.ocr'}:
                evidence=self._provider_operation(b,a,source)
            elif a['action'].startswith('excel.'):
                evidence = self._excel(b, a, source)
            elif a['action'].startswith('table.'):
                evidence = self._table(b, a, source)
            else:
                evidence = self._pdf(b, a, source)
            # Read results are bounded before entering the control-plane journal.
            # Artifact contents live in their file, not inline in OperationResult.
            if (a['action'].endswith('inspect') and
                    len(json.dumps(evidence, ensure_ascii=False, allow_nan=False)) > b.max_text_chars):
                raise ExecutorFailure('document_evidence_size_limit')
            evidence.setdefault('recomputed',False)
            return dict(evidence, input_sha256=original, providers=dependency_versions())

    def _provider_operation(self,b,a,source):
        recalculating=a['action']=='excel.recalculate'
        providers=b.recalculation_backends if recalculating else b.ocr_backends
        provider=next(p for p in providers if p.name==a['backend'])
        executable=provider.executable if recalculating else provider.tesseract_executable
        if executable is not None and not Path(executable).is_file():
            raise ExecutorFailure('configured_document_executable_missing')
        # Prepare a private input copy: providers can never save into the user's
        # original. Its temp folder is under the authorized output directory.
        output=Path(a['output']); source_digest=digest(source)
        with tempfile.TemporaryDirectory(prefix='.sentra-provider-',dir=output.parent) as temporary:
            directory=Path(temporary); staged=directory/'input.xlsx' if recalculating else directory/'input.pdf'
            effect_checkpoint(); shutil.copyfile(source,staged)
            if digest(staged)!=source_digest: raise ExecutorFailure('document_input_changed_while_copying')
            result_path=directory/('result.xlsx' if recalculating else 'result.json')
            config={'kind':provider.engine if recalculating else 'ocr','input':str(staged),'output':str(result_path),
                'workdir':str(directory),'executable':executable,'timeout':provider.timeout_seconds,'token':uuid.uuid4().hex}
            if recalculating:
                lib=dependency('openpyxl',(3,1,5))
                with zipfile.ZipFile(staged) as archive:
                    if sum(i.file_size for i in archive.infolist())>b.max_expanded_bytes:
                        raise ExecutorFailure('xlsx_expansion_limit')
                    if any(n=='xl/vbaProject.bin' or n.startswith('xl/externalLinks/') for n in archive.namelist()):
                        raise ExecutorFailure('recalculation_external_links_or_macros_denied')
                original_book=lib.load_workbook(staged,read_only=True,data_only=False,keep_links=False)
                try:
                    for ws in original_book.worksheets:
                        if ws.max_row>b.max_rows or ws.max_column>b.max_columns:
                            raise ExecutorFailure('worksheet_dimension_limit')
                    for ref in a['expected_cache']:
                        sheet,coordinate=ref.rsplit('!',1)
                        value=original_book[sheet][coordinate].value
                        if not isinstance(value,str) or not value.startswith('='):
                            raise ExecutorFailure('expected_cache_cell_not_formula')
                finally: original_book.close()
            else:
                pdf=dependency('pypdf',(6,10,0))
                with staged.open('rb') as stream:
                    reader=pdf.PdfReader(stream,strict=True)
                    if reader.is_encrypted: raise ExecutorFailure('encrypted_pdf_unsupported')
                    count=len(reader.pages)
                pages=a.get('pages',list(range(1,count+1)))
                if count>b.max_pages or not pages or len(pages)>b.max_pages or any(p>count for p in pages):
                    raise ExecutorFailure('pdf_page_out_of_range')
                config.update(pages=pages,language=provider.language,dpi=provider.dpi,
                    max_pixels=provider.max_pixels,max_text_bytes=b.max_text_chars)
            physical=run_owned_worker(provider.python_executable,DOCUMENT_WORKER_SCRIPT,config,
                                      timeout_seconds=provider.timeout_seconds)
            if not result_path.is_file() or result_path.stat().st_size>b.max_expanded_bytes:
                raise ExecutorFailure('document_provider_output_missing_or_oversized')
            def verify(path):
                if digest(source)!=source_digest:
                    raise ExecutorFailure('document_input_changed_before_publish')
                if recalculating:
                    lib=dependency('openpyxl',(3,1,5))
                    book=lib.load_workbook(path,read_only=True,data_only=True,keep_links=False)
                    try:
                        for ref,expected in a['expected_cache'].items():
                            sheet,coordinate=ref.rsplit('!',1)
                            actual=_json_cell(book[sheet][coordinate].value)
                            if type(actual) in (int,float) and type(expected) in (int,float):
                                matched=math.isclose(actual,expected,rel_tol=1e-12,abs_tol=1e-12)
                            else: matched=type(actual) is type(expected) and actual==expected
                            if not matched: raise ExecutorFailure('recalculated_formula_cache_verification_failed')
                    finally: book.close()
                else:
                    data=json.loads(path.read_text(encoding='utf-8'))
                    if (data.get('ocr_performed') is not True or
                        [p['page'] for p in data['pages']]!=config['pages']):
                        raise ExecutorFailure('ocr_output_verification_failed')
            verify(result_path)
            artifact=atomic_output(b.paths,a['output'],lambda p:shutil.copyfile(result_path,p),verify,
                overwrite=a.get('overwrite',False),max_output_bytes=b.max_expanded_bytes)
            return {'artifact':artifact,'provider':physical,'recomputed':recalculating,
                    'expected_cache_verified':recalculating,'ocr_performed':not recalculating,
                    'original_preserved':digest(source)==source_digest}

    def _excel(self, b, a, source):
        provider = dependency('openpyxl', (3, 1, 5))
        with zipfile.ZipFile(source) as archive:
            if sum(i.file_size for i in archive.infolist()) > b.max_expanded_bytes:
                raise ExecutorFailure('xlsx_expansion_limit')
            # openpyxl cannot preserve arbitrary package parts on a roundtrip.
            # Reject known unsupported workbook features before transformation.
            names = archive.namelist()
            if a['action'] == 'excel.transform' and any(
                    n == 'xl/vbaProject.bin' or n.startswith(('xl/slicers/', 'xl/slicerCaches/', 'xl/activeX/',
                                  'xl/ctrlProps/', 'xl/embeddings/', 'xl/externalLinks/')) for n in names):
                raise ExecutorFailure('xlsx_unsupported_preservation_feature')
        mode = a.get('value_mode', 'formulas')
        modifying = a['action'] == 'excel.transform'
        book = provider.load_workbook(source, read_only=not modifying,
                                      data_only=mode == 'cached_values', keep_links=False)
        try:
            sheet = book[a['sheet']] if 'sheet' in a else book.active
            if sheet.max_row > b.max_rows or sheet.max_column > b.max_columns:
                raise ExecutorFailure('worksheet_dimension_limit')
            if not modifying:
                start, end = a.get('start_row', 1), a.get('end_row', min(sheet.max_row, 100))
                if end < start:
                    raise ExecutorFailure('invalid_row_range')
                rows = [[_json_cell(c.value) for c in row] for row in
                        sheet.iter_rows(min_row=start, max_row=end)]
                return {'sheet': sheet.title, 'sheets': book.sheetnames,
                    'dimensions': [sheet.max_row, sheet.max_column],
                    'start_row': start, 'rows': rows, 'value_mode': mode,
                    'cached_values_may_be_missing_or_stale': mode == 'cached_values',
                    'truncated': start > 1 or end < sheet.max_row}
            for coordinate, value in a['cells'].items():
                cell = sheet[coordinate]
                if cell.row > b.max_rows or cell.column > b.max_columns:
                    raise ExecutorFailure('cell_outside_dimension_limit')
                cell.value = value
            from openpyxl.workbook.properties import CalcProperties
            book.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
            def verify(temp):
                check = provider.load_workbook(temp, data_only=False, read_only=True)
                try:
                    if check.sheetnames != book.sheetnames or any(
                            check[sheet.title][c].value != v for c, v in a['cells'].items()):
                        raise ExecutorFailure('xlsx_output_verification_failed')
                finally:
                    check.close()
            artifact = atomic_output(b.paths, a['output'], book.save, verify,
                                     overwrite=a.get('overwrite', False))
            return {'artifact': artifact, 'sheet': sheet.title, 'changed_cells': list(a['cells']),
                    'formula_cache_invalidated': True, 'recalculation_requested_on_open': True,
                    'preservation': 'openpyxl_supported_xlsx_features_only'}
        finally:
            book.close()

    def _table(self, b, a, source):
        delimiter = '\t' if source.suffix.lower() == '.tsv' else ','
        with source.open('r', encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream, delimiter=delimiter)
            columns = reader.fieldnames
            if (not columns or any(not c for c in columns) or
                    len(columns) != len(set(columns)) or len(columns) > b.max_columns):
                raise ExecutorFailure('invalid_table_header')
            rows = []
            for row in reader:
                if None in row or None in row.values():
                    raise ExecutorFailure('ragged_table_row')
                rows.append(row)
                if len(rows) > b.max_rows:
                    raise ExecutorFailure('table_row_limit')
        if a['action'] == 'table.inspect':
            limit = a.get('limit', min(100, b.max_rows))
            return {'columns': columns, 'row_count': len(rows), 'rows': rows[:limit],
                    'truncated': len(rows) > limit, 'value_types': 'UTF-8 strings'}
        columns, rows = transform_table(columns, rows, **{k: a[k] for k in (
            'conversions', 'filters', 'select', 'sort_by', 'group_by', 'sums') if k in a})
        target = Path(a['output'])
        if target.suffix.lower() == '.json':
            payload = json.dumps({'columns': columns, 'rows': rows}, ensure_ascii=False,
                                 allow_nan=False, indent=2).encode('utf-8')
            def verify(temp):
                if json.loads(temp.read_text(encoding='utf-8')) != {'columns': columns, 'rows': rows}:
                    raise ExecutorFailure('table_output_verification_failed')
        else:
            buffer = io.StringIO(newline='')
            delim = '\t' if target.suffix.lower() == '.tsv' else ','
            writer = csv.DictWriter(buffer, fieldnames=columns, delimiter=delim)
            writer.writeheader()
            writer.writerows(rows)
            payload = buffer.getvalue().encode('utf-8')
            expected = [{c: '' if row[c] is None else str(row[c]) for c in columns} for row in rows]
            def verify(temp):
                with temp.open(encoding='utf-8', newline='') as stream:
                    check = csv.DictReader(stream, delimiter=delim)
                    if check.fieldnames != columns or list(check) != expected:
                        raise ExecutorFailure('table_output_verification_failed')
        artifact = atomic_output(b.paths, a['output'], lambda p: p.write_bytes(payload),
                                 verify, overwrite=a.get('overwrite', False))
        return {'artifact': artifact, 'columns': columns, 'row_count': len(rows)}

    def _pdf(self, b, a, source):
        provider = dependency('pypdf', (6, 10, 0))
        # File context bounds ownership and closes handles even on parser failure.
        with source.open('rb') as stream:
            reader = provider.PdfReader(stream, strict=True)
            if reader.is_encrypted:
                raise ExecutorFailure('encrypted_pdf_unsupported')
            count = len(reader.pages)
            if count > b.max_pages:
                raise ExecutorFailure('pdf_page_limit')
            if a['action'] == 'pdf.inspect':
                return {'page_count': count, 'encrypted': False,
                    'pages': [{'page': i + 1, 'width': float(p.mediabox.width),
                               'height': float(p.mediabox.height)} for i, p in enumerate(reader.pages)]}
            pages = a.get('pages', list(range(1, count + 1)))
            if not pages or len(pages) > b.max_pages or any(p > count for p in pages):
                raise ExecutorFailure('pdf_page_out_of_range')
            if a['action'] == 'pdf.text':
                texts, remaining = [], b.max_text_chars
                truncated = False
                for number in pages:
                    page = reader.pages[number - 1]
                    content = page.get_contents()
                    if content and len(content.get_data()) > b.max_expanded_bytes:
                        raise ExecutorFailure('pdf_content_expansion_limit')
                    text = page.extract_text() or ''
                    texts.append({'page': number, 'text': text[:remaining]})
                    truncated |= len(text) > remaining
                    remaining = max(0, remaining - len(text))
                return {'page_count': count, 'pages': texts, 'truncated': truncated,
                        'ocr_performed': False, 'extraction': 'embedded PDF text; visual fidelity not asserted'}
            writer = provider.PdfWriter()
            try:
                for number in pages:
                    writer.add_page(reader.pages[number - 1])
                def produce(temp):
                    with temp.open('wb') as out:
                        writer.write(out)
                def verify(temp):
                    with temp.open('rb') as check_stream:
                        check = provider.PdfReader(check_stream, strict=True)
                        if len(check.pages) != len(pages):
                            raise ExecutorFailure('pdf_output_verification_failed')
                        for actual, number in zip(check.pages, pages):
                            expected = reader.pages[number - 1]
                            actual_content, expected_content = actual.get_contents(), expected.get_contents()
                            if (list(actual.mediabox) != list(expected.mediabox) or
                                (actual_content.get_data() if actual_content is not None else b'') !=
                                (expected_content.get_data() if expected_content is not None else b'')):
                                raise ExecutorFailure('pdf_output_verification_failed')
                artifact = atomic_output(b.paths, a['output'], produce, verify,
                                         overwrite=a.get('overwrite', False))
                return {'artifact': artifact, 'source_pages': pages, 'page_count': len(pages)}
            finally:
                writer.close()


def declare_document_machine(*, machine_id, owner_principal_id, bindings, policy):
    from sentra_runtime.contracts import Capability, Machine
    from .discovery import MachineDeclaration
    executor = DocumentExecutor(machine_id=machine_id, owner_principal_id=owner_principal_id,
                                bindings=bindings, policy=policy)
    return MachineDeclaration(Machine(machine_id, executor.kind, owner_principal_id,
        tuple(Capability(b.capability_id, 'Scoped deterministic documents', 'high')
              for b in bindings)), executor)
