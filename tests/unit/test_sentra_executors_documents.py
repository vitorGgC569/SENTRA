"""Real file acceptance, optional providers skip explicitly; no fake PDF/Excel SDK."""
import asyncio
import importlib.util
import json
import tempfile
import unittest
import zipfile
import os
import sys
from pathlib import Path

from sentra_executors.documents import (DocumentBinding, declare_document_machine, transform_table,
    compare_tables, FormulaRecalculationBinding, OCRBinding)
from sentra_executors.rpa import AuthorizedPaths, digest
from sentra_runtime.contracts import OperationRequest, PolicyDecision
from sentra_runtime.executor import ExecutorRegistry


class DocumentsAcceptance(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix='sentra-documents-acceptance-')
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name).resolve()
        self.source = self.root / 'source.csv'
        self.source.write_text('team,amount,name\nB,2,Beta\nA,10,Alpha\nB,3,Gamma\n', encoding='utf-8')
        self.counter = 0
        self.allowed = True
        self.policy = lambda r: PolicyDecision(self.allowed, 'acceptance file scope')
        self.binding = DocumentBinding('documents.files', AuthorizedPaths((str(self.root),), (str(self.root),)))
        self.declaration = declare_document_machine(machine_id='docs-acceptance',
            owner_principal_id='acceptance', bindings=(self.binding,), policy=self.policy)
        self.registry = ExecutorRegistry(authorize=self.policy)
        self.declaration.register(self.registry)

    def request(self, **arguments):
        self.counter += 1
        return OperationRequest(f'doc-{self.counter}', 'acceptance', 'docs-acceptance',
            'documents.files', 'file-acceptance', f'doc-key-{self.counter}', arguments)

    def invoke(self, **arguments):
        return asyncio.run(self.registry.submit(self.request(**arguments)))

    def test_csv_numeric_group_filter_and_verified_json(self):
        before = digest(self.source)
        output = self.root / 'result.json'
        result = self.invoke(action='table.transform', input=str(self.source), output=str(output),
            conversions={'amount': 'number'}, filters=[{'column': 'amount', 'op': 'gt', 'value': 1}],
            sort_by=['team'], group_by=['team'], sums=['amount'])
        self.assertEqual(result.state, 'SUCCEEDED', result.error)
        actual = json.loads(output.read_text(encoding='utf-8'))
        self.assertEqual(actual['rows'], [{'team': 'A', 'count': 1, 'sum_amount': 10.0},
                                         {'team': 'B', 'count': 2, 'sum_amount': 5.0}])
        self.assertEqual(digest(self.source), before)
        self.assertEqual(result.evidence['artifact']['sha256'], digest(output))
        self.assertTrue(result.evidence['artifact']['verified_before_publish'])
        self.assertFalse(result.evidence['recomputed'])

    def test_conversion_failure_preserves_existing_output_and_input(self):
        self.source.write_text('amount\nnot-a-number\n', encoding='utf-8')
        before = self.source.read_bytes()
        output = self.root / 'keep.csv'
        output.write_bytes(b'original destination')
        result = self.invoke(action='table.transform', input=str(self.source), output=str(output),
                             overwrite=True, conversions={'amount': 'number'})
        self.assertEqual((result.state, result.error), ('FAILED', 'numeric_conversion_failed'))
        self.assertEqual(output.read_bytes(), b'original destination')
        self.assertEqual(self.source.read_bytes(), before)
        self.assertFalse([p for p in self.root.glob('.sentra-*') if p.is_file()])

    def test_csv_tsv_roundtrip_projection(self):
        output = self.root / 'selected.tsv'
        result = self.invoke(action='table.transform', input=str(self.source), output=str(output),
                             select=['name', 'team'], sort_by=['name'])
        self.assertEqual(result.state, 'SUCCEEDED', result.error)
        observed = self.invoke(action='table.inspect', input=str(output), limit=2)
        self.assertEqual(observed.evidence['columns'], ['name', 'team'])
        self.assertEqual(observed.evidence['row_count'], 3)
        self.assertTrue(observed.evidence['truncated'])
        self.assertEqual(observed.evidence['rows'][0], {'name': 'Alpha', 'team': 'A'})

    def test_overwrite_requires_explicit_option(self):
        output = self.root / 'result.csv'
        output.write_bytes(b'keep')
        result = self.invoke(action='table.transform', input=str(self.source), output=str(output))
        self.assertEqual((result.state, result.error), ('FAILED', 'output_exists'))
        self.assertEqual(output.read_bytes(), b'keep')

    def test_input_cannot_be_output_even_when_explicit_overwrite(self):
        before = self.source.read_bytes()
        result = self.invoke(action='table.transform', input=str(self.source),
                             output=str(self.source), overwrite=True)
        self.assertEqual(result.state, 'FAILED')
        self.assertEqual(self.source.read_bytes(), before)

    def test_bad_scope_missing_file_and_unknown_fields_fail_closed(self):
        outside = self.root.parent / 'outside.csv'
        for arguments in (
            {'action': 'table.inspect', 'input': str(outside)},
            {'action': 'table.inspect', 'input': str(self.root / 'absent.csv')},
            {'action': 'table.inspect', 'input': str(self.source), 'evaluate': 'arbitrary code'},
        ):
            with self.subTest(arguments=arguments):
                result = asyncio.run(self.declaration.adapter.start(self.request(**arguments)))
                self.assertEqual(result.state, 'FAILED')
                self.assertEqual(result.error, 'invalid_scope_or_capability')

    def test_denied_policy_never_writes(self):
        self.allowed = False
        output = self.root / 'denied.csv'
        result = asyncio.run(self.declaration.adapter.start(self.request(action='table.transform',
            input=str(self.source), output=str(output))))
        self.assertEqual(result.error, 'policy_denied')
        self.assertFalse(output.exists())

    def test_idempotent_submission_does_not_repeat_publication(self):
        output = self.root / 'once.csv'
        request = self.request(action='table.transform', input=str(self.source), output=str(output))
        first = asyncio.run(self.registry.submit(request))
        second = asyncio.run(self.registry.submit(request))
        self.assertEqual(first.state, 'SUCCEEDED')
        self.assertEqual(first, second)

    def test_ragged_csv_is_not_silently_normalized(self):
        self.source.write_text('a,b\none,two,three\n', encoding='utf-8')
        result = self.invoke(action='table.inspect', input=str(self.source))
        self.assertEqual(result.error, 'ragged_table_row')

    def test_explicit_artifact_table_verifier_ignores_order_without_losing_duplicates(self):
        expected=self.root/'expected.csv'
        expected.write_text('team,amount,name\nB,3,Gamma\nB,2,Beta\nA,10,Alpha\n',encoding='utf-8')
        result=self.invoke(action='artifact.verify',input=str(self.source),expected=str(expected),
                           options={'ignore_row_order':True})
        self.assertEqual(result.state,'SUCCEEDED',result.error)
        self.assertTrue(result.evidence['passed'])
        self.assertEqual(result.evidence['expected_sha256'],digest(expected))
        expected.write_text(expected.read_text(encoding='utf-8')+'B,2,Beta\n',encoding='utf-8')
        result=self.invoke(action='artifact.verify',input=str(self.source),expected=str(expected),
                           options={'ignore_row_order':True})
        self.assertFalse(result.evidence['passed'])

    def test_composite_short_circuit_and_report_are_explicit(self):
        bad=self.root/'different.csv'; bad.write_text('team,amount,name\nX,0,Other\n',encoding='utf-8')
        output=self.root/'acceptance.json'
        result=self.invoke(action='acceptance.evaluate',output=str(output),conjunction='and',
            criteria=[{'id':'table','actual':str(self.source),'expected':str(bad)},
                      {'id':'bytes','actual':str(self.source),'expected':str(self.source),'metric':'bytes'}])
        self.assertEqual(result.state,'SUCCEEDED',result.error)  # Evaluation ran; acceptance failed.
        self.assertFalse(result.evidence['passed'])
        self.assertTrue(result.evidence['short_circuited'])
        self.assertEqual(result.evidence['criteria'][1]['state'],'NOT_EVALUATED')
        self.assertFalse(json.loads(output.read_text(encoding='utf-8'))['passed'])

    def test_expected_artifact_digest_pin_and_report_overwrite_protection(self):
        expected=self.root/'expected.csv'; expected.write_bytes(self.source.read_bytes())
        result=self.invoke(action='artifact.verify',input=str(self.source),expected=str(expected),expected_sha256='0'*64)
        self.assertEqual(result.error,'expected_artifact_digest_changed')
        before=digest(expected)
        result=self.invoke(action='artifact.verify',input=str(self.source),expected=str(expected),output=str(expected),overwrite=True)
        self.assertEqual(result.state,'FAILED')
        self.assertEqual(digest(expected),before)

    def test_formula_backend_requires_explicit_configuration(self):
        schema_only=self.root/'not-dispatched.xlsx'; schema_only.write_bytes(b'schema validation fixture; never parsed')
        result=self.invoke(action='excel.recalculate',input=str(schema_only),output=str(self.root/'result.xlsx'),
                           backend='implicit-excel',expected_cache={'Data!B1':10})
        self.assertEqual(result.state,'FAILED')

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'real openpyxl provider unavailable')
    def test_excel_formula_and_cache_are_distinct_without_recalculation(self):
        import openpyxl
        source = self.root / 'formulas.xlsx'
        book = openpyxl.Workbook()
        book.active.title = 'Data'
        book.active['A1'] = 2
        book.active['B1'] = '=A1*3'
        book.save(source)
        book.close()
        before = digest(source)
        formulas = self.invoke(action='excel.inspect', input=str(source), value_mode='formulas')
        cached = self.invoke(action='excel.inspect', input=str(source), value_mode='cached_values')
        self.assertEqual(formulas.state, 'SUCCEEDED', formulas.error)
        self.assertEqual(formulas.evidence['rows'][0], [2, '=A1*3'])
        self.assertEqual(cached.evidence['rows'][0], [2, None])
        self.assertTrue(cached.evidence['cached_values_may_be_missing_or_stale'])
        output = self.root / 'changed.xlsx'
        transformed = self.invoke(action='excel.transform', input=str(source), output=str(output),
                                 cells={'A1': 4, 'B1': '=A1*5'})
        self.assertEqual(transformed.state, 'SUCCEEDED', transformed.error)
        self.assertEqual(digest(source), before)
        self.assertFalse(transformed.evidence['recomputed'])
        actual = openpyxl.load_workbook(output, data_only=False)
        try:
            self.assertEqual(actual['Data']['B1'].value, '=A1*5')
            self.assertTrue(actual.calculation.fullCalcOnLoad)
        finally:
            actual.close()

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'real openpyxl provider unavailable')
    def test_excel_grid_verifier_does_not_require_tabular_headers(self):
        import openpyxl
        actual=self.root/'grid.xlsx'; expected=self.root/'expected-grid.xlsx'
        book=openpyxl.Workbook(); book.active['A1']=2; book.active['B1']='=A1*5'; book.save(actual); book.save(expected); book.close()
        result=self.invoke(action='artifact.verify',input=str(actual),expected=str(expected),metric='excel',
                           options={'compare_styles':True,'value_mode':'formulas'})
        self.assertEqual(result.state,'SUCCEEDED',result.error)
        self.assertTrue(result.evidence['passed']); self.assertFalse(result.evidence['recomputed'])

    @unittest.skipUnless(importlib.util.find_spec('pypdf'), 'real pypdf provider unavailable')
    def test_pdf_embedded_text_and_reordered_pages(self):
        import pypdf
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        source = self.root / 'source.pdf'
        writer = pypdf.PdfWriter()
        for width, text in ((200, 'First page'), (300, 'Second page')):
            page = writer.add_blank_page(width=width, height=400)
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
                DictionaryObject({NameObject('/F1'): font})})
            content = DecodedStreamObject()
            content.set_data(f'BT /F1 12 Tf 20 200 Td ({text}) Tj ET'.encode())
            page[NameObject('/Contents')] = content
        with source.open('wb') as stream:
            writer.write(stream)
        writer.close()
        before = digest(source)
        text = self.invoke(action='pdf.text', input=str(source), pages=[2])
        self.assertEqual(text.state, 'SUCCEEDED', text.error)
        self.assertIn('Second page', text.evidence['pages'][0]['text'])
        self.assertFalse(text.evidence['ocr_performed'])
        output = self.root / 'reordered.pdf'
        extracted = self.invoke(action='pdf.extract_pages', input=str(source),
                                output=str(output), pages=[2, 1, 2])
        self.assertEqual(extracted.state, 'SUCCEEDED', extracted.error)
        with output.open('rb') as stream:
            actual = pypdf.PdfReader(stream)
            self.assertEqual([float(p.mediabox.width) for p in actual.pages], [300, 200, 300])
            self.assertIn('Second page', actual.pages[0].extract_text())
        self.assertEqual(digest(source), before)

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'real openpyxl provider unavailable')
    def test_excel_unsupported_package_parts_are_not_silently_dropped(self):
        import openpyxl
        source = self.root / 'extended.xlsx'
        book = openpyxl.Workbook()
        book.active['A1'] = 'original'
        book.save(source)
        book.close()
        # An extra OOXML part exercises the preservation guard, not an Excel app.
        with zipfile.ZipFile(source, 'a') as archive:
            archive.writestr('xl/externalLinks/externalLink1.xml', '<externalLink/>')
        before = digest(source)
        output = self.root / 'protected.xlsx'
        output.write_bytes(b'existing output')
        result = self.invoke(action='excel.transform', input=str(source), output=str(output),
                             overwrite=True, cells={'A1': 'changed'})
        self.assertEqual(result.error, 'xlsx_unsupported_preservation_feature')
        self.assertEqual(output.read_bytes(), b'existing output')
        self.assertEqual(digest(source), before)

    @unittest.skipUnless(importlib.util.find_spec('pypdf'), 'real pypdf provider unavailable')
    def test_out_of_range_pdf_pages_leave_no_output(self):
        import pypdf
        source = self.root / 'one.pdf'
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=200, height=300)
        with source.open('wb') as stream:
            writer.write(stream)
        writer.close()
        output = self.root / 'denied.pdf'
        result = self.invoke(action='pdf.extract_pages', input=str(source), output=str(output), pages=[2])
        self.assertEqual(result.error, 'pdf_page_out_of_range')
        self.assertFalse(output.exists())


class PureTableContract(unittest.TestCase):
    def test_no_implicit_numeric_conversion(self):
        columns, rows = transform_table(['amount'], [{'amount': '10'}, {'amount': '2'}], sort_by=['amount'])
        self.assertEqual([r['amount'] for r in rows], ['10', '2'])
        self.assertEqual(columns, ['amount'])

    def test_numeric_tolerance_uses_explicit_columns_and_unique_keys(self):
        actual=[{'id':'b','amount':'2.001'},{'id':'a','amount':'1.001'}]
        expected=[{'id':'a','amount':'1'},{'id':'b','amount':'2'}]
        result=compare_tables(['id','amount'],actual,['id','amount'],expected,
            ignore_row_order=True,numeric_columns=['amount'],key_columns=['id'],absolute_tolerance=0.01)
        self.assertTrue(result['passed'])


@unittest.skipUnless(os.environ.get('SENTRA_RECALC_ACCEPTANCE')=='1','real Excel/LibreOffice acceptance staged, not enabled')
class RealRecalculationAcceptance(unittest.TestCase):
    def test_actual_formula_cache_is_recomputed_and_original_preserved(self):
        import openpyxl
        with tempfile.TemporaryDirectory(prefix='sentra-recalc-acceptance-') as folder:
            root=Path(folder).resolve(); source=root/'input.xlsx'; output=root/'result.xlsx'
            book=openpyxl.Workbook(); book.active.title='Data'; book.active['A1']=2; book.active['B1']='=A1*5'
            book.save(source); book.close(); before=digest(source)
            engine=os.environ.get('SENTRA_RECALC_ENGINE')
            self.assertIn(engine,('excel','libreoffice'),'configure the actual installed engine explicitly')
            provider=FormulaRecalculationBinding('real','acceptance',engine,
                python_executable=os.environ.get('SENTRA_RECALC_PYTHON',sys.executable),
                executable=os.environ.get('SENTRA_LIBREOFFICE_EXECUTABLE'),timeout_seconds=60)
            binding=DocumentBinding('recalc',AuthorizedPaths((str(root),),(str(root),)),
                                    timeout_seconds=90,recalculation_backends=(provider,))
            declaration=declare_document_machine(machine_id='recalc-acceptance',owner_principal_id='acceptance',
                bindings=(binding,),policy=lambda r:PolicyDecision(True,'real recalculation acceptance'))
            result=asyncio.run(declaration.adapter.start(OperationRequest('recalc','acceptance','recalc-acceptance',
                'recalc','work','recalc-key',{'action':'excel.recalculate','input':str(source),'output':str(output),
                                         'backend':'real','expected_cache':{'Data!B1':10}})))
            self.assertEqual(result.state,'SUCCEEDED',(result.error,result.evidence))
            self.assertTrue(result.evidence['recomputed']); self.assertEqual(digest(source),before)
            check=openpyxl.load_workbook(output,data_only=True)
            try:self.assertEqual(check['Data']['B1'].value,10)
            finally:check.close()


@unittest.skipUnless(os.environ.get('SENTRA_OCR_ACCEPTANCE')=='1','real Tesseract/PyMuPDF acceptance staged, not enabled')
class RealOCRAcceptance(unittest.TestCase):
    def test_scanned_pdf_uses_actual_render_and_tesseract(self):
        import pymupdf
        from PIL import Image,ImageDraw,ImageFont
        with tempfile.TemporaryDirectory(prefix='sentra-ocr-acceptance-') as folder:
            root=Path(folder).resolve(); source=root/'scan.pdf'; output=root/'ocr.json'
            image=Image.new('RGB',(1200,400),'white')
            font_path=os.environ.get('SENTRA_OCR_TEST_FONT','C:/Windows/Fonts/arial.ttf')
            font=ImageFont.truetype(font_path,60)
            ImageDraw.Draw(image).text((40,100),'SENTRA 12345',fill='black',font=font)
            image.save(root/'scan.png')
            pdf=pymupdf.open(); page=pdf.new_page(width=600,height=200)
            page.insert_image(page.rect,filename=str(root/'scan.png')); pdf.save(source); pdf.close()
            before=digest(source)
            executable=os.environ.get('SENTRA_TESSERACT_EXECUTABLE')
            self.assertTrue(executable,'configure the real installed Tesseract executable')
            provider=OCRBinding('real','acceptance',executable,
                python_executable=os.environ.get('SENTRA_OCR_PYTHON',sys.executable))
            binding=DocumentBinding('ocr',AuthorizedPaths((str(root),),(str(root),)),timeout_seconds=90,ocr_backends=(provider,))
            declaration=declare_document_machine(machine_id='ocr-acceptance',owner_principal_id='acceptance',
                bindings=(binding,),policy=lambda r:PolicyDecision(True,'real OCR acceptance'))
            result=asyncio.run(declaration.adapter.start(OperationRequest('ocr','acceptance','ocr-acceptance','ocr','work','ocr-key',
                {'action':'pdf.ocr','input':str(source),'output':str(output),'backend':'real','pages':[1]})))
            self.assertEqual(result.state,'SUCCEEDED',(result.error,result.evidence))
            actual=json.loads(output.read_text(encoding='utf-8'))
            self.assertTrue(actual['ocr_performed']); self.assertIn('12345',actual['pages'][0]['text'])
            self.assertEqual(digest(source),before)


if __name__ == '__main__':
    unittest.main()
